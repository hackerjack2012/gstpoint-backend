import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

def clean_sheet_name(name):
    invalid_chars = ['\\', '/', '*', '[', ']', ':', '?', "'"]
    for ch in invalid_chars:
        name = name.replace(ch, '')
    name = name.replace(' ', '_')
    return name[:31]

def multi_ledger_match(
    file_path,
    delay_threshold=180,
    gst_rate=5,
    output_path=None
):
    try:
        df = pd.read_excel(file_path, sheet_name="Sheet1", header=None)
        today = pd.to_datetime("today").normalize()

        supplier_blocks = []
        start_row = None
        supplier_name = None
        for i, row in df.iterrows():
            if (pd.notnull(row[0]) and
                (pd.isnull(row[1]) or row[1] == '') and
                (pd.isnull(row[2]) or row[2] == '') and
                str(row[0]).lower() != 'dates'):
                if start_row is not None:
                    supplier_blocks.append((supplier_name, start_row, i-1))
                supplier_name = str(row[0])
                start_row = i + 1
        if start_row is not None and supplier_name:
            supplier_blocks.append((supplier_name, start_row, len(df)-1))

        # Build a list of all supplier names (order preserved)
        all_suppliers = [block[0] for block in supplier_blocks]
        supplier_info_dict = {}

        wb = Workbook()
        summary_headers = ['Supplier Name', 'Delayed Rows', 'GST Rate (%)', 'Interest Total']
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

        # Process each supplier and collect info for summary
        for idx, (sup_name, start, end) in enumerate(supplier_blocks):
            block = df.iloc[start:end+1].copy()
            block.columns = ['Date', 'Debit', 'Credit']
            block = block.dropna(how='all')
            block = block[(block['Debit'].notna()) | (block['Credit'].notna())]
            block['Date'] = pd.to_datetime(block['Date'], errors='coerce')
            block['Debit'] = pd.to_numeric(block['Debit'], errors='coerce').fillna(0)
            block['Credit'] = pd.to_numeric(block['Credit'], errors='coerce').fillna(0)
            debits = block[block['Debit'] > 0].copy()
            credits = block[block['Credit'] > 0].copy()

            match_data = []
            delayed_rows = 0
            credit_balances = credits['Credit'].tolist()

            for i, debit_row in debits.iterrows():
                payment_date = debit_row['Date']
                payment_amt = debit_row['Debit']
                remaining_payment = payment_amt

                for j in range(len(credits)):
                    if credit_balances[j] > 0:
                        match_amt = min(remaining_payment, credit_balances[j])
                        delay_days = (payment_date - credits.iloc[j]['Date']).days
                        match_data.append({
                            'Purchase Date': credits.iloc[j]['Date'],
                            'Purchase Amount': match_amt,
                            'Payment Date': payment_date,
                            'Payment Used': match_amt,
                            'Days Delayed': delay_days,
                            'IsDelayed': delay_days > delay_threshold
                        })
                        if delay_days > delay_threshold:
                            delayed_rows += 1
                        credit_balances[j] -= match_amt
                        remaining_payment -= match_amt
                        if remaining_payment == 0:
                            break

            for j in range(len(credit_balances)):
                if credit_balances[j] > 0:
                    delay_days = (today - credits.iloc[j]['Date']).days
                    match_data.append({
                        'Purchase Date': credits.iloc[j]['Date'],
                        'Purchase Amount': credit_balances[j],
                        'Payment Date': 'Unpaid',
                        'Payment Used': 0,
                        'Days Delayed': delay_days,
                        'IsDelayed': delay_days > delay_threshold
                    })
                    if delay_days > delay_threshold:
                        delayed_rows += 1

            safe_name = clean_sheet_name(sup_name)
            interest_total_cell = None

            if delayed_rows > 0:
                ws = wb.create_sheet(title=safe_name)
                ws.append([
                    'Purchase Date', 'Purchase Amount', 'Payment Date', 'Payment Used', 'Days Delayed', 'Interest', 'IsDelayed'
                ])
                interest_rows = []
                # The summary row for this supplier is always idx+2 (1-based Excel, +1 for header)
                summary_gst_cell = f"Summary!$C${idx+2}"

                for n, row in enumerate(match_data):
                    ws.append([
                        row['Purchase Date'].strftime('%d-%m-%Y') if isinstance(row['Purchase Date'], pd.Timestamp) else row['Purchase Date'],
                        row['Purchase Amount'],
                        row['Payment Date'].strftime('%d-%m-%Y') if isinstance(row['Payment Date'], pd.Timestamp) else row['Payment Date'],
                        row['Payment Used'],
                        row['Days Delayed'],
                        0,  # Interest, will set formula below
                        int(row['IsDelayed'])
                    ])
                for r in range(2, ws.max_row+1):
                    days = ws.cell(row=r, column=5).value
                    is_delayed = ws.cell(row=r, column=7).value
                    if int(is_delayed) == 1:
                        for c in range(1, 8):
                            ws.cell(row=r, column=c).fill = yellow_fill
                        purchase_amt_cell = f"B{r}"
                        days_cell = f"E{r}"
                        # Use ABSOLUTE reference for GST rate
                        interest_formula = (
                            f"=IF({summary_gst_cell}=\"\",0,"
                            f"ROUND(({purchase_amt_cell}*{summary_gst_cell}/(100+{summary_gst_cell}))*0.18*{days_cell}/365,2))"
                        )
                        ws.cell(row=r, column=6).value = interest_formula
                        interest_rows.append(r)
                    else:
                        ws.cell(row=r, column=6).value = 0

                for col in ws.columns:
                    max_length = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        try:
                            cell_len = len(str(cell.value))
                            if cell_len > max_length:
                                max_length = cell_len
                        except:
                            pass
                    ws.column_dimensions[col_letter].width = max_length + 2

                if interest_rows:
                    total_row = ws.max_row + 1

                    # Write both label and formula on the SAME row
                    ws.cell(row=total_row, column=5).value = "Interest Total:"
                    ws.cell(row=total_row, column=6).value = f"=SUM(F2:F{total_row-1})"

                    interest_total_cell = f"F{total_row}"
                else:
                    interest_total_cell = None

                ws["H1"] = "Back to Summary"
                ws["H1"].hyperlink = "#Summary!A1"
                ws["H1"].style = "Hyperlink"
                ws.cell(row=1, column=1).style = "Hyperlink"

            # Always record supplier info for summary, even if delayed_rows=0
            supplier_info_dict[sup_name] = {
                "delayed_rows": delayed_rows,
                "safe_name": safe_name,
                "interest_total_cell": interest_total_cell if delayed_rows > 0 else None
            }

        # Remove default sheet if present and unused
        if 'Sheet' in wb.sheetnames and len(wb.sheetnames) > 1:
            wb.remove(wb['Sheet'])

        # Build summary_data for ALL suppliers, so row index matches supplier order
        summary_data = []
        for i, sup_name in enumerate(all_suppliers):
            info = supplier_info_dict.get(sup_name, {})
            delayed_rows = info.get("delayed_rows", 0)
            safe_name = info.get("safe_name", clean_sheet_name(sup_name))
            interest_total_cell = info.get("interest_total_cell", None)
            # Interest formula references supplier's own sheet/cell, else 0
            interest_formula = f"='{safe_name}'!{interest_total_cell}" if interest_total_cell else 0
            summary_data.append([sup_name, delayed_rows, "", interest_formula])

        ws_summary = wb.create_sheet(title="Summary", index=0)
        ws_summary.append(summary_headers)
        for i, row in enumerate(summary_data):
            ws_summary.append(row)
            sheet_name = clean_sheet_name(row[0])
            link = f"#{sheet_name}!A1"
            ws_summary.cell(row=i+2, column=1).hyperlink = link
            ws_summary.cell(row=i+2, column=1).style = "Hyperlink"

        for col in ws_summary.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    cell_len = len(str(cell.value))
                    if cell_len > max_length:
                        max_length = cell_len
                except:
                    pass
            ws_summary.column_dimensions[col_letter].width = max_length + 2

        if output_path is None:
            output_path = file_path.replace(
                ".xlsx",
                "_MultiMatched.xlsx"
            )

        wb.save(output_path)
        return output_path
    except Exception as e:
        raise Exception(str(e))