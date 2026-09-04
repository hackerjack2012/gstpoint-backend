import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime

def match_fifo(file_path, delay_threshold, gst_rate, output_path=None):
    try:

        gst_rate = float(gst_rate)

        if gst_rate > 1:
            gst_rate = gst_rate / 100
        df = pd.read_excel(file_path, sheet_name="Sheet1")
        today = pd.to_datetime("today").normalize()

        df['Date'] = pd.to_datetime(df['Date'], errors='coerce')
        df['Credit'] = pd.to_numeric(df['Credit'], errors='coerce')
        df['Debit'] = pd.to_numeric(df['Debit'], errors='coerce')
        df = df[df['Date'].notna()]

        purchases = df[df['Credit'].notna() & (df['Credit'] > 0)].copy()
        payments = df[df['Debit'].notna() & (df['Debit'] > 0)].copy()

        purchases.reset_index(drop=True, inplace=True)
        payments.reset_index(drop=True, inplace=True)

        match_data = []
        purchase_balances = purchases['Credit'].tolist()

        for i, pay_row in payments.iterrows():
            payment_date = pay_row['Date']
            payment_amt = pay_row['Debit']
            remaining_payment = payment_amt

            for j in range(len(purchases)):
                if purchase_balances[j] > 0:
                    match_amt = min(remaining_payment, purchase_balances[j])
                    if match_amt > 0:
                        match_data.append({
                            'Payment Date': payment_date,
                            'Payment Used': match_amt,
                            'Purchase Date': purchases.loc[j, 'Date'],
                            'Purchase Amount': match_amt,
                            'Days Delayed': (payment_date - purchases.loc[j, 'Date']).days
                        })
                        purchase_balances[j] -= match_amt
                        remaining_payment -= match_amt
                        if remaining_payment == 0:
                            break

            if remaining_payment > 0 and (today - payment_date).days > delay_threshold:
                match_data.append({
                    'Payment Date': payment_date,
                    'Payment Used': remaining_payment,
                    'Purchase Date': 'Unmatched',
                    'Purchase Amount': 0,
                    'Days Delayed': (today - payment_date).days
                })

        for j in range(len(purchase_balances)):
            if purchase_balances[j] > 0:
                match_data.append({
                    'Payment Date': 'Unpaid',
                    'Payment Used': 0,
                    'Purchase Date': purchases.loc[j, 'Date'],
                    'Purchase Amount': purchase_balances[j],
                    'Days Delayed': (today - purchases.loc[j, 'Date']).days
                })

        result_df = pd.DataFrame(match_data)

        if not result_df.empty:
            for col in ['Payment Date', 'Purchase Date']:
                if col in result_df.columns:
                    result_df[col] = result_df[col].apply(
                        lambda x: x.strftime('%d-%m-%Y') if isinstance(x, pd.Timestamp) else x
                    )

            def calc_taxable(purchase_amt):
                return round(purchase_amt / (1 + gst_rate), 2)

            def calc_gst(purchase_amt):
                return round(purchase_amt - calc_taxable(purchase_amt), 2)

            def calc_interest(gst_amt, days):
                if days is None or days <= 0:
                    return 0.00

                return round(gst_amt * 18 / 100 * days / 365, 2)

            result_df['Taxable Amount'] = result_df['Purchase Amount'].apply(calc_taxable)
            result_df['GST Tax'] = result_df['Purchase Amount'].apply(calc_gst)
            result_df['Interest'] = result_df.apply(
                lambda row: calc_interest(
                    row['GST Tax'],
                    row['Days Delayed']
                ),
                axis=1
            )

        if output_path is None:
            output_path = file_path.replace(".xlsx", "_MatchedFIFO.xlsx")
        result_df.to_excel(output_path, index=False)

        wb = load_workbook(output_path)
        ws = wb.active
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        for row in range(2, ws.max_row + 1):
            cell_value = ws.cell(row=row, column=5).value
            try:
                days = int(float(cell_value))
                if days > delay_threshold:
                    for col in range(1, 9):  # Apply red fill up to Interest column
                        ws.cell(row=row, column=col).fill = red_fill
            except (ValueError, TypeError):
                continue

        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)

            for cell in col:
                try:
                    cell_len = len(str(cell.value))
                    if cell_len > max_length:
                        max_length = cell_len
                except:
                    pass

            ws.column_dimensions[col_letter].width = max_length + 2

        wb.save(output_path)
        return output_path

    except Exception as e:
        raise Exception(str(e))