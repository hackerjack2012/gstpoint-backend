from openpyxl import (
    load_workbook,
    Workbook
)

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from openpyxl.utils import (
    get_column_letter
)


# ============================================================
# TAXPAYER OUTPUT COLUMNS
# ============================================================

TAXPAYER_COLUMNS = [

    "GSTIN",

    "Legal Name",

    "Trade Name",

    "Registration Date",

    "Constitution",

    "GSTIN Status",

    "Taxpayer Type",

    # ----------------------------------------
    # JURISDICTION
    # ----------------------------------------

    "Administrative Office (Main Dealing Office)",

    "Other Office",

    # ----------------------------------------
    # BUSINESS DETAILS
    # ----------------------------------------

    "Principal Place of Business",

    "Aadhaar Authenticated",

    "Aadhaar Authentication Date",

    "e-KYC Verified",

    "Nature of Core Business Activity",

    "Nature of Business Activities",

    "E-Invoice Status",

    "Field Visit Conducted",

    # ----------------------------------------
    # GOODS
    # ----------------------------------------

    "Goods HSN",

    "Goods Description",

    # ----------------------------------------
    # SERVICES
    # ----------------------------------------

    "Services HSN/SAC",

    "Services Description",

    # ----------------------------------------
    # PROCESSING
    # ----------------------------------------

    "Processing Status",

    "Error / Remarks"
]


# ============================================================
# RETURN STATUS COLUMNS
# ============================================================

RETURN_COLUMNS = [

    "GSTIN",

    "Legal Name",

    "Financial Year",

    "Month",

    "GSTR-1 Status",

    "GSTR-1 Filing Date",

    "GSTR-3B Status",

    "GSTR-3B Filing Date",

    "Processing Status",

    "Error / Remarks"
]


# ============================================================
# READ GSTIN INPUT
# ============================================================

def read_gstins(path):

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    sheet = workbook.active

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    workbook.close()

    if not rows:
        raise ValueError(
            "Excel file is empty."
        )

    header = [

        str(value).strip().upper()
        if value is not None
        else ""

        for value in rows[0]
    ]

    if "GSTIN" in header:

        gstin_index = (
            header.index(
                "GSTIN"
            )
        )

        data_rows = rows[1:]

    else:

        gstin_index = 0

        data_rows = rows

    gstins = []

    for row in data_rows:

        if (
            not row
            or gstin_index >= len(row)
        ):
            continue

        value = row[
            gstin_index
        ]

        if value is None:
            continue

        gstin = str(
            value
        ).strip().upper()

        if not gstin:
            continue

        gstins.append(
            gstin
        )

    if not gstins:

        raise ValueError(
            "No GSTINs found "
            "in Excel file."
        )

    return gstins


# ============================================================
# SAVE WORKBOOK
# ============================================================

def save_records(
    records,
    columns,
    output_path,
    sheet_name
):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = sheet_name

    # --------------------------------------------------------
    # COLORS
    # --------------------------------------------------------

    HEADER_BLUE = "1F4E78"

    ADMIN_GOLD = "FFF2CC"

    ADMIN_FONT = "9C6500"

    SUCCESS_GREEN = "E2F0D9"

    FAILED_RED = "FCE4D6"

    BORDER_COLOR = "D9E1F2"

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    for col_number, column in enumerate(
        columns,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=col_number,
            value=column
        )

        # Administrative Office gets special
        # highlighting because it is the
        # main dealing office.

        if column == (
            "Administrative Office "
            "(Main Dealing Office)"
        ):

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=ADMIN_GOLD
            )

            cell.font = Font(
                bold=True,
                color=ADMIN_FONT
            )

        else:

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=HEADER_BLUE
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # --------------------------------------------------------
    # DATA
    # --------------------------------------------------------

    for row_number, record in enumerate(
        records,
        start=2
    ):

        for col_number, column in enumerate(
            columns,
            start=1
        ):

            value = record.get(
                column,
                ""
            )

            cell = sheet.cell(
                row=row_number,
                column=col_number,
                value=value
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            # Highlight Administrative Office
            # down the entire column.

            if column == (
                "Administrative Office "
                "(Main Dealing Office)"
            ):

                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=ADMIN_GOLD
                )

                cell.font = Font(
                    bold=True
                )

        # ----------------------------------------
        # Processing status highlighting
        # ----------------------------------------

        if "Processing Status" in columns:

            status_column = (
                columns.index(
                    "Processing Status"
                ) + 1
            )

            status_cell = sheet.cell(
                row=row_number,
                column=status_column
            )

            if str(
                status_cell.value
            ).lower() == "success":

                status_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=SUCCESS_GREEN
                )

            elif str(
                status_cell.value
            ).lower() == "failed":

                status_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=FAILED_RED
                )

    # --------------------------------------------------------
    # BORDER
    # --------------------------------------------------------

    thin = Side(
        style="thin",
        color=BORDER_COLOR
    )

    for row in sheet.iter_rows():

        for cell in row:

            cell.border = Border(
                left=thin,
                right=thin,
                top=thin,
                bottom=thin
            )

    # --------------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------------

    sheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # FILTER
    # --------------------------------------------------------

    sheet.auto_filter.ref = (
        sheet.dimensions
    )

    # --------------------------------------------------------
    # COLUMN WIDTH
    # --------------------------------------------------------

    for column_number, column in enumerate(
        columns,
        start=1
    ):

        letter = get_column_letter(
            column_number
        )

        if column == (
            "Administrative Office "
            "(Main Dealing Office)"
        ):

            width = 65

        elif column == "Other Office":

            width = 65

        elif column == (
            "Principal Place of Business"
        ):

            width = 50

        elif column in [
            "Goods Description",
            "Services Description"
        ]:

            width = 55

        elif column == (
            "Nature of Business Activities"
        ):

            width = 40

        elif column == (
            "Error / Remarks"
        ):

            width = 45

        elif column in [
            "Legal Name",
            "Trade Name"
        ]:

            width = 30

        elif column == "GSTIN":

            width = 19

        elif column in [
            "Goods HSN",
            "Services HSN/SAC"
        ]:

            width = 24

        else:

            width = 20

        sheet.column_dimensions[
            letter
        ].width = width

    # --------------------------------------------------------
    # ROW HEIGHT
    # --------------------------------------------------------

    sheet.row_dimensions[
        1
    ].height = 42

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    workbook.save(
        output_path
    )


# ============================================================
# TAXPAYER WORKBOOK
# ============================================================

def save_taxpayer_records(
    records,
    output_path
):

    save_records(
        records,
        TAXPAYER_COLUMNS,
        output_path,
        "Taxpayer Details"
    )


# ============================================================
# RETURN WORKBOOK
# ============================================================

def save_return_records(
    records,
    output_path
):

    save_records(
        records,
        RETURN_COLUMNS,
        output_path,
        "Return Status"
    )